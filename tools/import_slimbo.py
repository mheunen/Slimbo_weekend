#!/usr/bin/env python3
"""
Zet het Slimbo-weekenddocument (Excel) om naar het importbestand van de app:
app/src/main/assets/seed_medael_2026.json

Gebruik:  python3 tools/import_slimbo.py <pad-naar-xlsx>

Bewust NIET overgenomen: achternamen, adressen, postcodes, nationaliteit en
geboortedata (tabbladen Gastenfiche en Verjaardagen). Die zijn niet nodig in de app.
"""
import json
import re
import sys
import unicodedata
from datetime import datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

import openpyxl

TZ = ZoneInfo("Europe/Brussels")
OUT = Path(__file__).resolve().parent.parent / "app/src/main/assets/seed_medael_2026.json"

# ---- Weekend: vrijdag 6 t/m maandag 9 november 2026 ----------------------------
FRI = datetime(2026, 11, 6, tzinfo=TZ)
DAYS = {"Vrijdag": 0, "Zaterdag": 1, "Zondag": 2, "Maandag": 3}


def at(day_offset: int, hour: int, minute: int = 0) -> int:
    """Tijdstip (epoch ms) op dag X van het weekend, lokale Belgische tijd."""
    base = FRI.date() + timedelta(days=day_offset)
    return int(datetime(base.year, base.month, base.day, hour, minute, tzinfo=TZ).timestamp() * 1000)


def slug(s: str) -> str:
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")


# ---- Namen gelijktrekken (spelling in het document wisselt) ----------------------
ALIASES = {
    "melisss": "Melissa", "melissa": "Melissa", "moris": "Moris", "desiree": "Desiree",
    "melanie": "Mélanie", "mélanie": "Mélanie", "dennis l": "Dennis", "dennis": "Dennis",
    "vriendin": "Vriendin van Wesley", "vriendin wesley": "Vriendin van Wesley",
    "curro": "Curro", "amra": "Amra",
}


def norm_name(n: str) -> str:
    n = n.strip()
    key = n.lower()
    if key in ALIASES:
        return ALIASES[key]
    return n[:1].upper() + n[1:]


def split_names(s) -> list[str]:
    if not s:
        return []
    parts = re.split(r"\s*(?:,|&|\ben\b)\s*", str(s))
    return [norm_name(p) for p in parts if p.strip()]


# ---- Spelgegevens die vaststaan (min, max, minuten, zwaarte) ---------------------
# Alleen titels waarvan dit zeker is; de rest krijgt "nog controleren" in de app.
KNOWN = {
    "dice-hospital": (1, 4, 75, "MIDDEL"),
    "jaipur": (2, 2, 30, "LICHT"),
    "letter-jam": (2, 6, 45, "LICHT"),
    "cat-in-a-box": (2, 5, 30, "LICHT"),
    "curious-cargo": (2, 2, 45, "MIDDEL"),
    "sankore": (1, 4, 120, "ZWAAR"),
    "similo-slimbo": (2, 8, 15, "LICHT"),
    "trio": (3, 6, 15, "LICHT"),
    "bitoku": (1, 4, 120, "ZWAAR"),
    "galactic-cruise": (1, 4, 120, "ZWAAR"),
    "woodcraft": (1, 4, 90, "MIDDEL"),
    "a-gentle-rain": (1, 1, 20, "LICHT"),
    "maracaibo-en-uitbreiding-the-uprising": (1, 4, 120, "ZWAAR"),
    "nippon": (2, 4, 120, "ZWAAR"),
    "ohanami": (2, 4, 20, "LICHT"),
    "civolution": (1, 4, 180, "ZWAAR"),
    "inventions": (1, 4, 120, "ZWAAR"),
    "honey-buzz": (1, 4, 90, "MIDDEL"),
    "dice-forge": (2, 4, 45, "LICHT"),
    "arcs": (2, 4, 120, "ZWAAR"),
    "seti-evt-met-space-agencies": (1, 4, 160, "ZWAAR"),
    "endeavor-deep-sea": (1, 4, 90, "MIDDEL"),
    "the-gang": (3, 6, 20, "LICHT"),
    "shackleton-base": (1, 4, 120, "ZWAAR"),
    "arborea": (1, 4, 120, "ZWAAR"),
    "calimala": (1, 5, 90, "MIDDEL"),
    "century-spice-road": (2, 5, 45, "LICHT"),
    "container": (3, 5, 90, "MIDDEL"),
    "daitoshi": (1, 4, 120, "ZWAAR"),
    "dinosaur-island": (1, 4, 120, "MIDDEL"),
    "grand-austria-hotel": (2, 4, 120, "ZWAAR"),
    "het-witte-kasteel-het-duel": (2, 2, 30, "MIDDEL"),
    "hitster-alle-versies": (2, 10, 30, "LICHT"),
    "libertalia-winds-of-galecrest": (1, 6, 60, "MIDDEL"),
    "merlin-evt-met-uitbreiding-morgana": (2, 4, 90, "MIDDEL"),
    "parks": (1, 5, 60, "LICHT"),
    "sabika": (2, 4, 90, "ZWAAR"),
    "scholars-of-the-south-tigris": (1, 4, 120, "ZWAAR"),
    "tenpenny-parks": (1, 4, 75, "MIDDEL"),
    "the-castles-of-burgundy": (2, 4, 90, "MIDDEL"),
    "the-great-library": (1, 4, 150, "ZWAAR"),
}


def cell_text(v) -> str:
    return "" if v is None else str(v).strip()


def main(xlsx: str):
    wb = openpyxl.load_workbook(xlsx, data_only=True)

    # ---------------- Locaties & tafels (Molen van Medael) ----------------
    places, tables = [], []
    for nr in ("17", "18"):
        pid = f"res{nr}"
        places.append({
            "id": pid,
            "name": f"Residentie {nr}",
            "description": "Begane grond: woon/eetkamer met open keuken, terras. "
                           "1e en 2e verdieping: elk 4 slaapkamers en 4 badkamers.",
        })
        tables += [
            {"id": f"t{nr}-eet-a", "name": "Eettafel A", "placeId": pid, "seats": 11},
            {"id": f"t{nr}-eet-b", "name": "Eettafel B", "placeId": pid, "seats": 11},
            {"id": f"t{nr}-lounge", "name": "Loungehoek (salontafels)", "placeId": pid, "seats": 4},
            {"id": f"t{nr}-terras-1", "name": "Terras picknicktafel 1", "placeId": pid, "seats": 6},
            {"id": f"t{nr}-terras-2", "name": "Terras picknicktafel 2", "placeId": pid, "seats": 6},
        ]

    # ---------------- Kamers + gasten (tabblad Sleep arrangement) ----------------
    ws = wb["Sleep arrangement"]
    rooms, guests = [], {}
    current = None
    counter: dict[tuple, int] = {}
    for r in ws.iter_rows(values_only=True):
        apt, floor, kind, people = (cell_text(x) for x in r[:4])
        if apt.startswith("Residentie"):
            current = apt.split()[-1]
        if not current or not floor or not kind:
            continue
        try:
            fl = int(float(floor))
        except ValueError:
            continue
        beds = int(re.search(r"\d", kind).group())
        key = (current, fl)
        counter[key] = counter.get(key, 0) + 1
        idx = counter[key]
        rid = f"r{current}-{fl}-{idx}"
        names = split_names(people)
        rooms.append({
            "id": rid,
            "placeId": f"res{current}",
            "floor": fl,
            "name": f"Kamer {fl}.{idx}",
            "beds": beds,
            "bedInfo": f"{beds}× eenpersoons boxspring 90×200 (hoeslaken 90×200)",
            "occupants": names,
        })
        for n in names:
            guests[n] = {"id": "gst-" + slug(n), "name": n, "roomId": rid, "wants": []}

    # ---------------- Spellen (tabblad Games) ----------------
    ws = wb["Games"]
    games = []
    for r in list(ws.iter_rows(min_row=2, values_only=True)):
        name = cell_text(r[0])
        if not name:
            continue
        owner_raw = cell_text(r[1])
        wish = owner_raw.lower() == "wens"
        owner = "" if wish else norm_name(owner_raw)
        gid = "g-" + slug(name)
        link = cell_text(r[3])
        if not link.startswith("http"):
            link = "https://boardgamegeek.com/geeksearch.php?action=search&objecttype=boardgame&q=" + \
                   re.sub(r"\s+", "+", re.sub(r"\(.*?\)", "", name).strip())
        video = cell_text(r[4]) if cell_text(r[4]).startswith("http") else ""
        interested = split_names(r[2])
        known = KNOWN.get(slug(name))
        g = {
            "id": gid,
            "name": name,
            "owner": owner,
            "wishlist": wish,
            "link": link,
            "video": video,
            "interested": interested,
            "minPlayers": known[0] if known else 2,
            "maxPlayers": known[1] if known else 4,
            "durationMin": known[2] if known else 60,
            "weight": known[3] if known else "MIDDEL",
            "infoChecked": bool(known),
            "notes": "Wens: nog niemand neemt dit mee" if wish else "",
        }
        games.append(g)
        for n in interested:
            if n not in guests:
                guests[n] = {"id": "gst-" + slug(n), "name": n, "roomId": "", "wants": []}
            guests[n]["wants"].append(gid)

    # ---------------- Corvee (tabblad Corvee) ----------------
    ws = wb["Corvee"]
    chores = []
    slots = {"Ochtend": ("ochtend", 9, 13), "Middag": ("middag", 13, 18), "Avond": ("avond", 18, 22)}
    res = None
    slot = None
    for r in ws.iter_rows(values_only=True):
        a, b = cell_text(r[0]), cell_text(r[1])
        if a.startswith("Residentie"):
            res = a.split()[-1]
            continue
        if not res:
            continue
        for k in slots:
            if b.startswith(k):
                slot = k
        if not slot:
            continue
        for col, day in ((2, "Vrijdag"), (3, "Zaterdag"), (4, "Zondag"), (5, "Maandag")):
            v = cell_text(r[col]) if col < len(r) else ""
            if not v:
                continue
            key, h1, h2 = slots[slot]
            cid = f"c{res}-{slug(day)}-{key}"
            existing = next((c for c in chores if c["id"] == cid), None)
            if existing is None:
                existing = {
                    "id": cid, "placeId": f"res{res}", "day": day, "slot": key,
                    "startAt": at(DAYS[day], h1), "endAt": at(DAYS[day], h2), "names": [],
                }
                chores.append(existing)
            existing["names"] += split_names(v)

    # ---------------- Weekendinfo ----------------
    info = [
        {"title": "Aankomst en vertrek", "items": [
            "Inchecken vrijdag vanaf 16:00",
            "Uitchecken maandag uiterlijk 10:00 (maandag ochtend is er extra corvee)",
            "Parkeren bij de accommodatie, maar weinig plekken: carpool!",
            "Elektrisch laden kan niet bij de accommodatie zelf (volgens het landgoed wel een laadpunt aan de overkant)",
        ]},
        {"title": "Zelf meenemen", "items": [
            "Dekbedovertrek en kussensloop (dekbed ligt er)",
            "Hoeslaken 90×200 (eenpersoons boxsprings)",
            "2 theedoeken per persoon — zonder theedoek geen entree",
            "Handdoeken en douchespullen",
            "Eigen eten en drinken, met naamsticker",
            "Identiteitsbewijs",
        ]},
        {"title": "Gezamenlijk ingekocht (€1–2 p.p., verrekend met de borg)", "items": [
            "Wc-papier, keukenrol, vuilniszakken + PMD- en GFT-zakken",
            "Zonnebloemolie, vloeibare boter, peper en zout",
            "Afwasmiddel, borstel, spons, vaatwastabletten",
            "Zeep en handdoekjes voor de gezamenlijke wc's",
            "Curry, ketchup, mayo (ook vegan), zoetjes, bakpapier",
            "Douchegordijn + stang voor de kamers zonder deur",
        ]},
        {"title": "Afspraken", "items": [
            "Alles wat je meebrengt, gaat ook weer mee naar huis",
            "Koelkast: één eruit, één erin — niet alles tegelijk koud zetten (2 koelkasten per residentie)",
            "Ruim je eigen vaat op en laat de keuken netjes achter; pannen was je zelf af",
            "Vaatwasser vol? Aanzetten of uitruimen",
            "Corvee is een extra taak voor de gezamenlijke ruimtes",
            "Lichte slaper? Neem oordopjes mee, het kan gehorig zijn",
            "Kamers ruilen mag, geef het door aan Wilma",
        ]},
        {"title": "Murder Mystery", "items": [
            "Er kan weer een Murder Mystery georganiseerd worden: max. 20 personen, bijdrage €2–3. Aanmelden via het document.",
        ]},
    ]

    ws = wb["Carpool"]
    carpool = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not cell_text(r[0]) or cell_text(r[0]).lower() == "naam":
            continue
        parts = [f"{cell_text(r[0])} ({cell_text(r[1])})", f"heen {cell_text(r[2])}", f"terug {cell_text(r[3])}"]
        if cell_text(r[4]).lower() == "ja":
            plek = cell_text(r[5]).replace(".0", "")
            parts.append(f"{plek} plek")
        if cell_text(r[6]):
            parts.append("rijdt mee: " + cell_text(r[6]))
        carpool.append(" · ".join(parts))
    if carpool:
        info.append({"title": "Carpool", "items": carpool})

    ws = wb["Te koop aangeboden"]
    sale = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not cell_text(r[0]):
            continue
        price = r[2]
        price_s = f"€{price:g}" if isinstance(price, (int, float)) else cell_text(price)
        extra = " — " + cell_text(r[3]) if cell_text(r[3]) else ""
        sale.append(f"{cell_text(r[0])} · {cell_text(r[1])} · {price_s}{extra}")
    if sale:
        info.append({"title": "Te koop aangeboden", "items": sale})

    activities = [
        {"id": "a-inchecken", "type": "ANDERS", "title": "Inchecken Molen van Medael",
         "startAt": at(0, 16), "durationMin": 60, "placeId": "", "notes": "Vanaf 16:00"},
        {"id": "a-uitchecken", "type": "ANDERS", "title": "Uitchecken (uiterlijk 10:00)",
         "startAt": at(3, 9, 30), "durationMin": 30, "placeId": "", "notes": "Corvee maandag ochtend: zie rooster"},
        {"id": "a-diner-vrijdag", "type": "ETEN", "title": "Avondeten: saté van Desirée",
         "startAt": at(0, 18, 30), "durationMin": 90, "placeId": "",
         "notes": "Verzorgd door Desirée. Tijd is een schatting; pas aan indien nodig."},
        {"id": "a-diner-zaterdag", "type": "ETEN", "title": "Avondeten: Limburgs buffet",
         "startAt": at(1, 18, 30), "durationMin": 90, "placeId": "",
         "notes": "Tijd is een schatting; pas aan indien nodig."},
    ]

    seed = {
        "version": 1,
        "event": {
            "venue": "Landgoed Molen van Medael",
            "address": "Rue de Rémersdael 98, 4852 Hombourg (België)",
            "phone": "+32 87 78 43 50",
            "checkInAt": at(0, 16),
            "checkOutAt": at(3, 10),
            "info": info,
        },
        "places": places,
        "tables": tables,
        "rooms": rooms,
        "guests": sorted(guests.values(), key=lambda g: g["name"].lower()),
        "games": games,
        "chores": chores,
        "activities": activities,
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(seed, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"{OUT}: {len(games)} spellen ({sum(g['infoChecked'] for g in games)} met bekende spelersaantallen), "
          f"{len(rooms)} kamers, {len(guests)} gasten, {len(chores)} corveeblokken")


if __name__ == "__main__":
    main(sys.argv[1])
